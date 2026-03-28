#!/usr/bin/env python3
"""LLM-guided template section scanner and extractor for bid documents.

This script keeps deterministic operations in code while leaving semantic
boundary choices to the LLM/operator.
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import subprocess
import tempfile
from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
from pathlib import Path
from typing import Iterable

from docx import Document

DEFAULT_KEYWORDS = [
    "响应文件格式",
    "投标文件格式",
    "应答文件格式",
    "谈判应答文件格式",
    "比选文件格式",
    "文件格式",
    "附件",
]

CHAPTER_RE = re.compile(r"^\s*第[一二三四五六七八九十百零〇0-9]+[章节部分编]")


def normalize_text(text: str) -> str:
    text = text.replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def paragraph_text(el) -> str:
    text = "".join((node.text or "") for node in el.iter() if node.tag.endswith("}t"))
    return normalize_text(text)


def is_paragraph(el) -> bool:
    return el.tag.endswith("}p")


def is_chapter_heading(text: str, strict: bool = False) -> bool:
    """Return whether text looks like a top-level chapter heading."""
    if not CHAPTER_RE.match(text):
        return False
    if not strict:
        return len(text) <= 90

    if len(text) > 50:
        return False
    if any(ch in text for ch in "，。；：！？,.!?"):
        return False
    return True


def run_soffice_convert(
    input_path: Path, outdir: Path, profile_dir: Path, target_format: str
) -> set[Path]:
    before = set(outdir.glob(f"*.{target_format}"))
    cmd = [
        "soffice",
        "-env:UserInstallation=file://" + str(profile_dir),
        "--invisible",
        "--headless",
        "--norestore",
        "--convert-to",
        target_format,
        "--outdir",
        str(outdir),
        str(input_path),
    ]
    subprocess.run(cmd, check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    after = set(outdir.glob(f"*.{target_format}"))
    return after - before


def pick_latest(paths: set[Path]) -> Path | None:
    if not paths:
        return None
    return sorted(paths, key=lambda p: p.stat().st_mtime, reverse=True)[0]


def load_as_docx(input_path: Path, temp_root: Path) -> Path:
    suffix = input_path.suffix.lower()
    if suffix == ".docx":
        return input_path
    if suffix != ".doc":
        raise ValueError(f"Unsupported extension: {input_path.suffix}. Use .docx, .doc, or .pdf")

    outdir = temp_root / "converted"
    outdir.mkdir(parents=True, exist_ok=True)
    profile_dir = temp_root / "lo_profile"
    profile_dir.mkdir(parents=True, exist_ok=True)

    created_docx = run_soffice_convert(input_path, outdir, profile_dir, "docx")
    docx = pick_latest(created_docx)
    if docx:
        return docx

    # Fallback: convert .doc -> .odt -> .docx for documents that fail direct conversion.
    odt_created = run_soffice_convert(input_path, outdir, profile_dir, "odt")
    odt = pick_latest(odt_created)
    if odt:
        created_retry = run_soffice_convert(odt, outdir, profile_dir, "docx")
        retry_docx = pick_latest(created_retry)
        if retry_docx:
            return retry_docx

    candidates = sorted(outdir.glob("*.docx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if candidates:
        return candidates[0]

    raise RuntimeError(f"Failed to convert .doc to .docx: {input_path}")


def convert_pdf_with_soffice(input_path: Path, temp_root: Path) -> Path | None:
    outdir = temp_root / "converted_pdf_soffice"
    outdir.mkdir(parents=True, exist_ok=True)
    profile_dir = temp_root / "lo_profile_pdf"
    profile_dir.mkdir(parents=True, exist_ok=True)

    created = run_soffice_convert(input_path, outdir, profile_dir, "docx")
    docx = pick_latest(created)
    if docx:
        return docx

    candidates = sorted(outdir.glob("*.docx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if candidates:
        return candidates[0]
    return None


def convert_pdf_with_pdf2docx(input_path: Path, temp_root: Path) -> Path | None:
    try:
        from pdf2docx import Converter  # type: ignore
    except Exception:
        return None

    outdir = temp_root / "converted_pdf_pdf2docx"
    outdir.mkdir(parents=True, exist_ok=True)
    out_path = outdir / f"{input_path.stem}.docx"
    conv = Converter(str(input_path))
    devnull = open(os.devnull, "w", encoding="utf-8")
    prev_disable = logging.root.manager.disable
    try:
        logging.disable(logging.CRITICAL)
        with redirect_stdout(devnull), redirect_stderr(devnull):
            conv.convert(str(out_path))
    finally:
        logging.disable(prev_disable)
        devnull.close()
        conv.close()
    if out_path.exists() and out_path.stat().st_size > 0:
        return out_path
    return None


def load_pdf_as_docx(
    input_path: Path, temp_root: Path, preferred: str = "pdf2docx"
) -> tuple[Path | None, str | None]:
    if preferred not in {"pdf2docx", "soffice"}:
        raise ValueError("preferred must be 'pdf2docx' or 'soffice'")

    if preferred == "pdf2docx":
        first = [("pdf_pdf2docx", convert_pdf_with_pdf2docx), ("pdf_soffice_docx", convert_pdf_with_soffice)]
    else:
        first = [("pdf_soffice_docx", convert_pdf_with_soffice), ("pdf_pdf2docx", convert_pdf_with_pdf2docx)]

    for mode, fn in first:
        converted = fn(input_path, temp_root)
        if converted:
            return converted, mode
    return None, None


def load_pdf_lines(input_path: Path, temp_root: Path) -> list[str]:
    txt_path = temp_root / f"{input_path.stem}.txt"
    cmd = [
        "pdftotext",
        "-layout",
        "-enc",
        "UTF-8",
        str(input_path),
        str(txt_path),
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    text = txt_path.read_text(encoding="utf-8", errors="ignore")
    lines = [normalize_text(line) for line in text.splitlines()]
    rows = [line for line in lines if line]
    if not rows:
        raise RuntimeError(
            f"No extractable text found in PDF: {input_path}. "
            "The file may be image-only and require OCR."
        )
    return rows


def load_pdf_pages_lines(input_path: Path, temp_root: Path) -> list[list[str]]:
    txt_path = temp_root / f"{input_path.stem}.pages.txt"
    cmd = [
        "pdftotext",
        "-layout",
        "-enc",
        "UTF-8",
        str(input_path),
        str(txt_path),
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    text = txt_path.read_text(encoding="utf-8", errors="ignore")

    raw_pages = text.split("\f")
    pages: list[list[str]] = []
    for page in raw_pages:
        lines = [normalize_text(line) for line in page.splitlines()]
        lines = [line for line in lines if line]
        pages.append(lines)

    while pages and not pages[-1]:
        pages.pop()
    if not pages:
        raise RuntimeError(
            f"No extractable text found in PDF: {input_path}. "
            "The file may be image-only and require OCR."
        )
    return pages


def iter_body_paragraphs(doc: Document):
    para_idx = 0
    body = list(doc._element.body.iterchildren())
    for body_idx, el in enumerate(body):
        if not is_paragraph(el):
            continue
        text = paragraph_text(el)
        if not text:
            continue
        para_idx += 1
        yield para_idx, body_idx, text, el


def iter_pdf_lines(lines: list[str]):
    for para_idx, text in enumerate(lines, start=1):
        yield para_idx, para_idx - 1, text


def iter_pdf_page_lines(pages: list[list[str]]):
    para_idx = 0
    body_idx = 0
    for page_no, lines in enumerate(pages, start=1):
        for line_no_in_page, text in enumerate(lines, start=1):
            para_idx += 1
            yield para_idx, body_idx, text, page_no, line_no_in_page
            body_idx += 1


def heading_match(text: str, target: str, mode: str) -> bool:
    text_norm = normalize_text(text)
    target_norm = normalize_text(target)
    if mode == "exact":
        return text_norm == target_norm
    return target_norm in text_norm


def find_heading_occurrence(doc: Document, heading: str, occurrence: int, match_mode: str):
    hits = []
    for para_idx, body_idx, text, _ in iter_body_paragraphs(doc):
        if heading_match(text, heading, match_mode):
            hits.append((para_idx, body_idx, text))
    if occurrence <= 0:
        raise ValueError("occurrence must be >= 1")
    if len(hits) < occurrence:
        raise RuntimeError(
            f"Heading not found at occurrence {occurrence}: {heading}. Hits={len(hits)}"
        )
    return hits[occurrence - 1], hits


def guess_end_by_next_chapter(doc: Document, start_body_idx: int) -> tuple[int, str | None]:
    body = list(doc._element.body.iterchildren())
    for body_idx in range(start_body_idx + 1, len(body)):
        el = body[body_idx]
        if not is_paragraph(el):
            continue
        text = paragraph_text(el)
        if not text:
            continue
        if is_chapter_heading(text, strict=True):
            return body_idx, text
    return len(body), None


def guess_end_by_next_chapter_lines(lines: list[str], start_body_idx: int) -> tuple[int, str | None]:
    for body_idx in range(start_body_idx + 1, len(lines)):
        text = lines[body_idx]
        if is_chapter_heading(text, strict=True):
            return body_idx, text
    return len(lines), None


def clear_doc_body(out_doc: Document) -> None:
    out_body = out_doc._element.body
    sect = out_body.sectPr
    for child in list(out_body):
        if child is not sect:
            out_body.remove(child)


def copy_body_slice(doc: Document, start_body_idx: int, end_body_idx: int, output_path: Path) -> dict:
    body = list(doc._element.body.iterchildren())

    out_doc = Document()
    clear_doc_body(out_doc)
    out_body = out_doc._element.body
    sect = out_body.sectPr

    copied = 0
    copied_paragraphs = 0
    copied_tables = 0
    for el in body[start_body_idx:end_body_idx]:
        if el.tag.endswith("}sectPr"):
            continue
        sect.addprevious(deepcopy(el))
        copied += 1
        if el.tag.endswith("}p"):
            copied_paragraphs += 1
        elif el.tag.endswith("}tbl"):
            copied_tables += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    out_doc.save(str(output_path))

    return {
        "copied_blocks": copied,
        "copied_paragraphs": copied_paragraphs,
        "copied_tables": copied_tables,
    }


def copy_line_slice_to_docx(
    lines: list[str], start_body_idx: int, end_body_idx: int, output_path: Path
) -> dict:
    out_doc = Document()
    clear_doc_body(out_doc)
    for text in lines[start_body_idx:end_body_idx]:
        out_doc.add_paragraph(text)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    out_doc.save(str(output_path))
    copied = max(0, end_body_idx - start_body_idx)
    return {
        "copied_blocks": copied,
        "copied_paragraphs": copied,
        "copied_tables": 0,
    }


def copy_pdf_pages_to_docx(
    input_pdf: Path, start_page: int, end_page_inclusive: int, output_path: Path
) -> dict:
    if start_page <= 0:
        raise ValueError("start_page must be >= 1")
    if end_page_inclusive < start_page:
        raise ValueError("end_page_inclusive must be >= start_page")

    with tempfile.TemporaryDirectory(prefix="bid_pdf_pages_") as td:
        out_prefix = Path(td) / "page"
        cmd = [
            "pdftoppm",
            "-png",
            "-f",
            str(start_page),
            "-l",
            str(end_page_inclusive),
            str(input_pdf),
            str(out_prefix),
        ]
        subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

        page_imgs = sorted(
            Path(td).glob("page-*.png"),
            key=lambda p: int(re.search(r"-(\d+)\.png$", p.name).group(1))
            if re.search(r"-(\d+)\.png$", p.name)
            else 0,
        )
        if not page_imgs:
            raise RuntimeError("Failed to render PDF pages to images")

        out_doc = Document()
        clear_doc_body(out_doc)
        section = out_doc.sections[0]
        page_width = section.page_width - section.left_margin - section.right_margin

        for idx, img in enumerate(page_imgs):
            para = out_doc.add_paragraph()
            run = para.add_run()
            run.add_picture(str(img), width=page_width)
            if idx < len(page_imgs) - 1:
                out_doc.add_page_break()

        output_path.parent.mkdir(parents=True, exist_ok=True)
        out_doc.save(str(output_path))

    copied_pages = len(page_imgs)
    return {
        "copied_blocks": copied_pages,
        "copied_paragraphs": copied_pages,
        "copied_tables": 0,
    }


def scan_command(input_path: Path, keywords: Iterable[str], pdf_converter: str) -> list[dict]:
    rows = []
    suffix = input_path.suffix.lower()
    with tempfile.TemporaryDirectory(prefix="bid_extract_") as td:
        if suffix == ".pdf":
            pdf_docx, _ = load_pdf_as_docx(input_path, Path(td), preferred=pdf_converter)
            if pdf_docx:
                doc = Document(str(pdf_docx))
                for para_idx, body_idx, text, _ in iter_body_paragraphs(doc):
                    reasons = []
                    if is_chapter_heading(text, strict=False):
                        reasons.append("chapter")
                    for kw in keywords:
                        if kw and kw in text:
                            reasons.append(f"kw:{kw}")
                    if reasons:
                        rows.append(
                            {
                                "para_index": para_idx,
                                "body_index": body_idx,
                                "reason": ",".join(reasons),
                                "text": text,
                            }
                        )
                return rows

            lines = load_pdf_lines(input_path, Path(td))
            for para_idx, body_idx, text in iter_pdf_lines(lines):
                reasons = []
                if is_chapter_heading(text, strict=False):
                    reasons.append("chapter")
                for kw in keywords:
                    if kw and kw in text:
                        reasons.append(f"kw:{kw}")
                if reasons:
                    rows.append(
                        {
                            "para_index": para_idx,
                            "body_index": body_idx,
                            "reason": ",".join(reasons),
                            "text": text,
                        }
                    )
            return rows

        docx_path = load_as_docx(input_path, Path(td))
        doc = Document(str(docx_path))

        for para_idx, body_idx, text, _ in iter_body_paragraphs(doc):
            reasons = []
            if is_chapter_heading(text, strict=False):
                reasons.append("chapter")
            for kw in keywords:
                if kw and kw in text:
                    reasons.append(f"kw:{kw}")
            if reasons:
                rows.append(
                    {
                        "para_index": para_idx,
                        "body_index": body_idx,
                        "reason": ",".join(reasons),
                        "text": text,
                    }
                )
    return rows


def write_markdown_report(path: Path, summary: dict) -> None:
    lines = [
        "# Template Extraction Report",
        "",
        f"- input: `{summary['input']}`",
        f"- output: `{summary['output']}`",
        f"- source_mode: `{summary.get('source_mode', 'docx')}`",
        f"- pdf_converter: `{summary.get('pdf_converter') or ''}`",
        f"- start_heading: `{summary['start_heading']}`",
        f"- start_occurrence: `{summary['start_occurrence']}`",
        f"- start_para_index: `{summary['start_para_index']}`",
        f"- start_body_index: `{summary['start_body_index']}`",
    ]
    if summary.get("end_heading"):
        lines.extend(
            [
                f"- end_heading: `{summary['end_heading']}`",
                f"- end_occurrence: `{summary['end_occurrence']}`",
                f"- end_para_index: `{summary.get('end_para_index')}`",
                f"- end_body_index: `{summary['end_body_index']}`",
            ]
        )
    else:
        lines.extend(
            [
                "- end_heading: `AUTO_NEXT_CHAPTER`",
                f"- next_chapter_heading: `{summary.get('next_chapter_heading') or ''}`",
                f"- end_body_index: `{summary['end_body_index']}`",
            ]
        )
    if summary.get("start_page") is not None:
        lines.append(f"- start_page: `{summary['start_page']}`")
    if summary.get("end_page") is not None:
        lines.append(f"- end_page: `{summary['end_page']}`")
    lines.extend(
        [
            f"- copied_blocks: `{summary['copied_blocks']}`",
            f"- copied_paragraphs: `{summary['copied_paragraphs']}`",
            f"- copied_tables: `{summary['copied_tables']}`",
            "",
        ]
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Scan and extract bid template sections.")
    sub = parser.add_subparsers(dest="command", required=True)

    scan = sub.add_parser("scan", help="Print candidate markers for LLM boundary selection")
    scan.add_argument("input", type=Path, help="Input .docx, .doc, or .pdf")
    scan.add_argument(
        "--keywords",
        default=",".join(DEFAULT_KEYWORDS),
        help="Comma-separated keywords to mark in scan output",
    )
    scan.add_argument(
        "--pdf-converter",
        choices=["pdf2docx", "soffice"],
        default="pdf2docx",
        help="Preferred converter for PDF scan path",
    )
    scan.add_argument("--json", action="store_true", help="Print JSON array")

    extract = sub.add_parser("extract", help="Extract template section by heading boundaries")
    extract.add_argument("input", type=Path, help="Input .docx, .doc, or .pdf")
    extract.add_argument("output", type=Path, help="Output .docx")
    extract.add_argument("--start-heading", required=True, help="Section start heading")
    extract.add_argument("--start-occurrence", type=int, default=1)
    extract.add_argument("--end-heading", help="Optional explicit section end heading")
    extract.add_argument("--end-occurrence", type=int, default=1)
    extract.add_argument(
        "--match-mode",
        choices=["exact", "contains"],
        default="exact",
        help="Heading matching mode",
    )
    extract.add_argument(
        "--pdf-converter",
        choices=["pdf2docx", "soffice"],
        default="pdf2docx",
        help="Preferred converter for PDF extract path",
    )
    extract.add_argument(
        "--pdf-fallback-mode",
        choices=["image_pages", "text"],
        default="text",
        help=(
            "When PDF cannot be converted to DOCX structurally, fallback to either "
            "high-fidelity page images or plain text paragraphs."
        ),
    )
    end_group = extract.add_mutually_exclusive_group()
    end_group.add_argument(
        "--stop-at-next-chapter",
        dest="stop_at_next_chapter",
        action="store_true",
        help="If end heading is missing, stop at next chapter heading",
    )
    end_group.add_argument(
        "--no-stop-at-next-chapter",
        dest="stop_at_next_chapter",
        action="store_false",
        help="If end heading is missing, keep copying until document end",
    )
    extract.set_defaults(stop_at_next_chapter=True)
    extract.add_argument("--report", type=Path, help="Optional markdown report path")
    extract.add_argument("--json", action="store_true", help="Print JSON summary")

    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    if args.command == "scan":
        keywords = [normalize_text(x) for x in args.keywords.split(",") if normalize_text(x)]
        rows = scan_command(args.input, keywords, args.pdf_converter)
        if args.json:
            print(json.dumps(rows, ensure_ascii=False, indent=2))
            return

        print("para_index\tbody_index\treason\ttext")
        for row in rows:
            print(
                f"{row['para_index']}\t{row['body_index']}\t{row['reason']}\t{row['text']}"
            )
        return

    with tempfile.TemporaryDirectory(prefix="bid_extract_") as td:
        suffix = args.input.suffix.lower()
        source_mode = "docx"
        start_page = None
        end_page = None
        if suffix == ".pdf":
            pdf_docx, pdf_docx_mode = load_pdf_as_docx(
                args.input, Path(td), preferred=args.pdf_converter
            )
            if pdf_docx:
                source_mode = str(pdf_docx_mode or "pdf_docx")
                doc = Document(str(pdf_docx))

                start_hit, _ = find_heading_occurrence(
                    doc=doc,
                    heading=args.start_heading,
                    occurrence=args.start_occurrence,
                    match_mode=args.match_mode,
                )
                start_para_idx, start_body_idx, start_text = start_hit

                next_chapter_heading = None
                end_para_idx = None
                if args.end_heading:
                    all_hits = []
                    for para_idx, body_idx, text, _ in iter_body_paragraphs(doc):
                        if body_idx <= start_body_idx:
                            continue
                        if heading_match(text, args.end_heading, args.match_mode):
                            all_hits.append((para_idx, body_idx, text))
                    if len(all_hits) < args.end_occurrence:
                        raise RuntimeError(
                            f"End heading not found at occurrence {args.end_occurrence}: {args.end_heading}. Hits={len(all_hits)}"
                        )
                    end_para_idx, end_body_idx, _ = all_hits[args.end_occurrence - 1]
                else:
                    if args.stop_at_next_chapter:
                        end_body_idx, next_chapter_heading = guess_end_by_next_chapter(
                            doc, start_body_idx
                        )
                    else:
                        end_body_idx = len(list(doc._element.body.iterchildren()))

                copied_stats = copy_body_slice(doc, start_body_idx, end_body_idx, args.output)
            else:
                pdf_pages = load_pdf_pages_lines(args.input, Path(td))
                flat_rows = list(iter_pdf_page_lines(pdf_pages))

                hits = []
                for para_idx, body_idx, text, page_no, line_no in flat_rows:
                    if heading_match(text, args.start_heading, args.match_mode):
                        hits.append((para_idx, body_idx, text, page_no, line_no))
                if len(hits) < args.start_occurrence:
                    raise RuntimeError(
                        f"Heading not found at occurrence {args.start_occurrence}: {args.start_heading}. Hits={len(hits)}"
                    )
                start_para_idx, start_body_idx, start_text, start_page, _ = hits[
                    args.start_occurrence - 1
                ]

                next_chapter_heading = None
                end_para_idx = None
                end_line_no = None
                if args.end_heading:
                    all_hits = []
                    for para_idx, body_idx, text, page_no, line_no in flat_rows:
                        if body_idx <= start_body_idx:
                            continue
                        if heading_match(text, args.end_heading, args.match_mode):
                            all_hits.append((para_idx, body_idx, text, page_no, line_no))
                    if len(all_hits) < args.end_occurrence:
                        raise RuntimeError(
                            f"End heading not found at occurrence {args.end_occurrence}: {args.end_heading}. Hits={len(all_hits)}"
                        )
                    end_para_idx, end_body_idx, _, end_page, end_line_no = all_hits[
                        args.end_occurrence - 1
                    ]
                else:
                    if args.stop_at_next_chapter:
                        end_body_idx = len(flat_rows)
                        for para_idx, body_idx, text, page_no, line_no in flat_rows:
                            if body_idx <= start_body_idx:
                                continue
                            if is_chapter_heading(text, strict=True):
                                end_body_idx = body_idx
                                end_para_idx = para_idx
                                end_page = page_no
                                end_line_no = line_no
                                next_chapter_heading = text
                                break
                    else:
                        end_body_idx = len(flat_rows)

                if args.pdf_fallback_mode == "image_pages":
                    source_mode = "pdf_page_images"
                    total_pages = len(pdf_pages)
                    if end_body_idx >= len(flat_rows):
                        render_end_page = total_pages
                    else:
                        render_end_page = int(end_page or total_pages)
                        if end_line_no is not None and end_line_no <= 2 and render_end_page > start_page:
                            render_end_page -= 1
                        if render_end_page < start_page:
                            render_end_page = start_page

                    copied_stats = copy_pdf_pages_to_docx(
                        args.input,
                        start_page=start_page,
                        end_page_inclusive=render_end_page,
                        output_path=args.output,
                    )
                    end_page = render_end_page
                else:
                    source_mode = "pdf_text_fallback"
                    lines = [text for _, _, text, _, _ in flat_rows]
                    copied_stats = copy_line_slice_to_docx(
                        lines, start_body_idx, end_body_idx, args.output
                    )
        else:
            input_docx = load_as_docx(args.input, Path(td))
            doc = Document(str(input_docx))

            start_hit, _ = find_heading_occurrence(
                doc=doc,
                heading=args.start_heading,
                occurrence=args.start_occurrence,
                match_mode=args.match_mode,
            )
            start_para_idx, start_body_idx, start_text = start_hit

            next_chapter_heading = None
            end_para_idx = None
            if args.end_heading:
                all_hits = []
                for para_idx, body_idx, text, _ in iter_body_paragraphs(doc):
                    if body_idx <= start_body_idx:
                        continue
                    if heading_match(text, args.end_heading, args.match_mode):
                        all_hits.append((para_idx, body_idx, text))
                if len(all_hits) < args.end_occurrence:
                    raise RuntimeError(
                        f"End heading not found at occurrence {args.end_occurrence}: {args.end_heading}. Hits={len(all_hits)}"
                    )
                end_para_idx, end_body_idx, _ = all_hits[args.end_occurrence - 1]
            else:
                if args.stop_at_next_chapter:
                    end_body_idx, next_chapter_heading = guess_end_by_next_chapter(doc, start_body_idx)
                else:
                    end_body_idx = len(list(doc._element.body.iterchildren()))

            copied_stats = copy_body_slice(doc, start_body_idx, end_body_idx, args.output)

        summary = {
            "input": str(args.input),
            "output": str(args.output),
            "pdf_converter": args.pdf_converter if suffix == ".pdf" else None,
            "start_heading": start_text,
            "start_occurrence": args.start_occurrence,
            "start_para_index": start_para_idx,
            "start_body_index": start_body_idx,
            "end_heading": args.end_heading,
            "end_occurrence": args.end_occurrence if args.end_heading else None,
            "end_para_index": end_para_idx,
            "end_body_index": end_body_idx,
            "next_chapter_heading": next_chapter_heading,
            "source_mode": source_mode,
            "start_page": start_page,
            "end_page": end_page,
            **copied_stats,
        }

        if args.report:
            write_markdown_report(args.report, summary)

        if args.json:
            print(json.dumps(summary, ensure_ascii=False, indent=2))
            return

        print("Extraction completed")
        print(f"- output: {summary['output']}")
        print(f"- start: para={summary['start_para_index']} body={summary['start_body_index']} text={summary['start_heading']}")
        if summary["end_heading"]:
            print(f"- end: para={summary['end_para_index']} body={summary['end_body_index']} text={summary['end_heading']}")
        else:
            print(f"- end: body={summary['end_body_index']} next_chapter={summary.get('next_chapter_heading') or 'N/A'}")
        print(f"- copied: blocks={summary['copied_blocks']} paragraphs={summary['copied_paragraphs']} tables={summary['copied_tables']}")


if __name__ == "__main__":
    main()
